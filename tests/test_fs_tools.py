from pathlib import Path

from docx import Document

from fs_tools import (
    list_files,
    read_file,
    search_in_file,
    write_file,
)


def test_read_txt_file(tmp_path):
    file_path = tmp_path / "resume.txt"
    file_path.write_text("Python\nFastAPI\nDocker", encoding="utf-8")

    result = read_file(str(file_path))

    assert result["success"] is True
    assert "Python" in result["content"]
    assert result["metadata"]["filename"] == "resume.txt"
    assert result["error"] is None


def test_read_missing_file():
    result = read_file("missing.txt")

    assert result["success"] is False
    assert result["error"] == "File not found."


def test_read_docx_file(tmp_path):
    file_path = tmp_path / "resume.docx"

    document = Document()
    document.add_paragraph("Python")
    document.add_paragraph("FastAPI")
    document.save(file_path)

    result = read_file(str(file_path))

    assert result["success"] is True
    assert "Python" in result["content"]


def test_list_all_files(tmp_path):
    (tmp_path / "a.txt").write_text("A", encoding="utf-8")
    (tmp_path / "b.pdf").write_text("PDF", encoding="utf-8")
    (tmp_path / "c.docx").write_text("DOCX", encoding="utf-8")

    files = list_files(str(tmp_path))

    assert len(files) == 3


def test_list_files_with_extension(tmp_path):
    (tmp_path / "one.txt").write_text("1", encoding="utf-8")
    (tmp_path / "two.txt").write_text("2", encoding="utf-8")
    (tmp_path / "three.pdf").write_text("3", encoding="utf-8")

    files = list_files(str(tmp_path), ".txt")

    assert len(files) == 2

    for file in files:
        assert file["name"].endswith(".txt")


def test_list_invalid_directory():
    files = list_files("does_not_exist")

    assert files == []


def test_write_file(tmp_path):
    destination = tmp_path / "nested" / "notes.txt"

    result = write_file(str(destination), "Hello World")

    assert result["success"] is True
    assert destination.exists()
    assert destination.read_text(encoding="utf-8") == "Hello World"


def test_search_keyword(tmp_path):
    file_path = tmp_path / "resume.txt"

    file_path.write_text(
        "Python\nFastAPI\nDocker\nPython Developer",
        encoding="utf-8",
    )

    result = search_in_file(str(file_path), "python")

    assert result["success"] is True
    assert result["count"] == 2
    assert result["matches"][0]["line"] == 1
    assert result["matches"][1]["line"] == 4


def test_search_no_match(tmp_path):
    file_path = tmp_path / "resume.txt"

    file_path.write_text(
        "Java\nSpring\nHibernate",
        encoding="utf-8",
    )

    result = search_in_file(str(file_path), "python")

    assert result["success"] is True
    assert result["count"] == 0
    assert result["matches"] == []


def test_search_missing_file():
    result = search_in_file("missing.txt", "python")

    assert result["success"] is False
    assert result["count"] == 0
    assert result["matches"] == []